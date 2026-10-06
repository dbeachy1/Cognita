"""Provider-free Cognita 16 audiobook production walkthrough.

Run from the repository root in the installed service environment:

    python examples/audiobook/synthetic_16_walkthrough.py

For checkout-only development, run with ``PYTHONPATH=src``. The installed
package run must report the site-packages module location at the release gate.

The example requires Cognita's registered FFmpeg and ffprobe paths (the
container defaults on Linux). It writes only to temporary directories. Its
small PCM samples and provider receipts are synthetic; it makes no network or
TTS calls and never reads a personal manuscript.
"""

from __future__ import annotations

import hashlib
import base64
import gc
import json
import math
import struct
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
import cognita

from cognita.books.backup_fixture import (
    restore_quiesced_book_project,
    snapshot_quiesced_book_project,
)
from cognita.books.fingerprint import canonical_json_sha256
from cognita.books.models import (
    BuildRequest,
    CommitBuildRequest,
    GetBookRequest,
    GetChapterRequest,
    GetGenerationsRequest,
    GetJobRequest,
    FindChunkRequest,
    ImportAudioRequest,
    InspectRequest,
    PrepareRequest,
    RecordGenerationRequest,
)
from cognita.books.projection import project_docx_pair
from cognita.books.service import BookService
from cognita.books.state import ProjectState
from cognita.books.policy import EffectiveIndexPolicy
from cognita.books.config import FolderRule
from cognita.books.configuration import load_book_config
from cognita.deindexed import DeindexedPaths
from cognita.books.policy import BookMutationPolicy
from cognita.config import CognitaConfig

PROJECT = "synthetic-16"
OWNER = "principal:synthetic-example"
JOB_IDS: list[str] = []
REQUEST_SPEC = {
    "provider": "synthetic-local-fixture",
    "route": "offline-example",
    "model_id": "deterministic-samples-v1",
    "voice_id": "fixture-voice",
    "parameters": {"fixture": True},
    "context_fields": {"language": "en"},
}
TARGET = {
    "sample_rate_hz": 44_100,
    "channels": 1,
    "encoding": "signed_integer",
    "storage_bits": 16,
    "valid_bits": 16,
    "mp3_bitrate_kbps": 192,
}


def write_docx(path: Path, text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    document.save(path)
    return path.read_bytes()


def write_tagged_docx(path: Path, text: str, spoken_tag: str) -> bytes:
    document = Document()
    paragraph = document.add_paragraph()
    split = text.index(".") + 1
    paragraph.add_run(text[:split])
    tag_style = document.styles.add_style("CognitaAudioTag", WD_STYLE_TYPE.CHARACTER)
    paragraph.add_run(" " + spoken_tag).style = tag_style
    paragraph.add_run(text[split:])
    document.save(path)
    return path.read_bytes()


def fixture_project(root: Path) -> dict[str, bytes]:
    now = datetime.now(timezone.utc).isoformat()
    (root / "Project Files/Source").mkdir(parents=True)
    (root / "Project Files/Source/approved-source.docx").write_bytes(
        write_docx(root / "Project Files/Source/source-staging.docx",
                   "Synthetic offline audiobook source."))
    (root / "Project Files/Source/source-staging.docx").unlink()
    chapter_ids = ["chapter-one", "chapter-two"]
    texts = {
        "chapter-one": "The copper door opened. The lantern stayed lit.",
        "chapter-two": "A small bell rang. The road continued west.",
    }
    chapters = []
    chapter_bytes = {}
    for index, chapter_id in enumerate(chapter_ids, start=1):
        relative = Path("Chapters") / str(index)
        directory = root / relative
        directory.mkdir(parents=True)
        prose_path = directory / "chapter.docx"
        tagged_path = directory / "chapter_audio-tags.docx"
        prose = write_docx(prose_path, texts[chapter_id])
        tagged = write_tagged_docx(tagged_path, texts[chapter_id], "softly")
        chapter_bytes[chapter_id] = prose
        (directory / "chapter.json").write_text(json.dumps({
            "schema_version": 1,
            "chapter_id": chapter_id,
            "layout_revision": 1,
            "state_revision": 1,
            "editorial_status": "approved",
            "approval_binding": "prose_projection",
            "approved_source_raw_sha256": hashlib.sha256(prose).hexdigest(),
            "approved_prose_projection_sha256": None,
            "approval_projection_version": None,
            "approval_provenance": None,
            "summary": None,
            "index_annotations": None,
        }), encoding="utf-8")
        projected = project_docx_pair(prose, tagged)
        state_path = directory / "chapter.json"
        chapter_state = json.loads(state_path.read_text(encoding="utf-8"))
        chapter_state.update({
            "approved_prose_projection_sha256": projected.prose_projection_sha256,
            "approval_projection_version": projected.projection_version,
            "approval_provenance": {
                "actor": "synthetic-example",
                "approved_at": now,
                "source_raw_sha256": projected.prose_sha256,
                "prose_projection_sha256": projected.prose_projection_sha256,
                "projection_version": projected.projection_version,
            },
        })
        state_path.write_text(json.dumps(chapter_state), encoding="utf-8")
        chapters.append({
            "chapter_id": chapter_id,
            "title": f"Chapter {index}",
            "chapter_state_filepath": (relative / "chapter.json").as_posix(),
            "working_filepath": (relative / "chapter.docx").as_posix(),
            "tagged_filepath": (relative / "chapter_audio-tags.docx").as_posix(),
            "summary_filepath": None,
            "originals_root": (relative / "Originals").as_posix(),
            "audio_root": f"Audiobook/Chapters/{index}",
        })
    layout = {
        "schema_version": 1,
        "layout_revision": 1,
        "book_id": "synthetic-book-16",
        "title": "Synthetic Offline Book",
        "chapter_order": chapter_ids,
        "chapters": chapters,
        "indexed_references": [],
        "indexed_instructions": [],
        "indexed_workflow_documents": [],
        "index_policy": {
            "default_unknown": "exclude", "tagged_copies": "exclude",
            "archives": "exclude", "media": "exclude",
            "duplicate_prose": "single_active_source",
        },
        "source_master_filepath": "Project Files/Source/approved-source.docx",
        "shared_paths": {
            "production_settings_filepath": "Project Files/production-settings.json",
            "book_audio_root": "Audiobook",
        },
        "storage": {
            "quota_bytes": 1_000_000_000,
            "reserve_bytes": 10_000_000,
            "import_https_hosts": [],
        },
        "production_authorization": {
            "authorization_id": "synthetic-production-approval",
            "actor": "synthetic-example",
            "authorized_at": now,
            "completed_book": True,
            "revoked": False,
        },
        "test_authorizations": [],
    }
    (root / "Project Files/Book_Layout.json").write_text(json.dumps(layout), encoding="utf-8")
    settings = {
        "schema_version": 1,
        "target_codepoints": 1000,
        "request_limit": {
            "value": 1000,
            "unit": "unicode_codepoints",
            "evidence": "synthetic offline example",
        },
        "request_spec": REQUEST_SPEC,
        "production_target": TARGET,
        "native_format_evidence": "Deterministic local PCM fixture; no provider used.",
    }
    (root / "Project Files/production-settings.json").write_text(
        json.dumps(settings, separators=(",", ":")), encoding="utf-8")
    return chapter_bytes


def prepare_chapter(service: BookService, chapter_id: str, prose: bytes,
                    operation: str, spoken_tag: str,
                    expected_manifest_revision: int | None = None) -> dict:
    index = 1 if chapter_id == "chapter-one" else 2
    tagged_path = service.root / f"Chapters/{index}/chapter_audio-tags.docx"
    tagged = write_tagged_docx(tagged_path, {
        "chapter-one": "The copper door opened. The lantern stayed lit.",
        "chapter-two": "A small bell rang. The road continued west.",
    }[chapter_id], spoken_tag)
    inspected = service.inspect(InspectRequest.model_validate({
        "project": PROJECT,
        "chapter_id": chapter_id,
        "prose_filepath": f"Chapters/{index}/chapter.docx",
        "tagged_filepath": f"Chapters/{index}/chapter_audio-tags.docx",
    }))
    second_sentence = " The lantern" if chapter_id == "chapter-one" else " The road"
    split = inspected["speech_text"].index(second_sentence)
    settings_path = service.root / "Project Files/production-settings.json"
    settings_bytes = settings_path.read_bytes()
    request = PrepareRequest.model_validate({
        "project": PROJECT,
        "operation_id": operation,
        "chapter_id": chapter_id,
        "document_view_id": inspected["document_view_id"],
        "expected_prose_sha256": hashlib.sha256(prose).hexdigest(),
        "expected_tagged_sha256": hashlib.sha256(tagged).hexdigest(),
        "expected_manifest_revision": expected_manifest_revision,
        "scope": {"kind": "production"},
        "speech_selection_confirmed": True,
        "request_limit": {"value": 1000, "unit": "unicode_codepoints"},
        "expected_settings_sha256": hashlib.sha256(settings_bytes).hexdigest(),
        "production_target": TARGET,
        "chunks": [
            {"chunk_id": f"{chapter_id}-opening", "start": 0, "end": split,
             "request_spec": REQUEST_SPEC},
            {"chunk_id": f"{chapter_id}-closing", "start": split,
             "end": len(inspected["speech_text"]), "request_spec": REQUEST_SPEC},
        ],
        "publish_bookmarks_to_working_tagged_docx": False,
    })
    return service.prepare(request, owner_key=OWNER)[0]


def synthetic_pcm(phase: int, frames: int = 44_100) -> bytes:
    values = bytearray()
    for frame in range(frames):
        sample = int(1800 * math.sin(2 * math.pi * 220 * frame / 44_100 + phase))
        values.extend(struct.pack("<h", sample))
    return bytes(values)


def make_take(service: BookService, prepared: dict, chapter_id: str,
              chunk_id: str, operation: str, phase: int,
              frames: int = 44_100) -> dict:
    chunk = next(item for item in prepared["chunks"] if item["chunk_id"] == chunk_id)
    reserved, replayed = service.record_generation(RecordGenerationRequest.model_validate({
        "project": PROJECT,
        "operation_id": f"{operation}-reserve",
        "change": {
            "kind": "reserve", "chapter_id": chapter_id,
            "snapshot_id": prepared["snapshot_id"], "chunk_id": chunk_id,
            "expected_manifest_revision": prepared["manifest_revision"],
            "request": {"prompt_sha256": chunk["prompt_sha256"], "spec": chunk["request_spec"]},
        },
    }), owner_key=OWNER)
    assert not replayed
    generation_id = reserved["generation"]["generation_record_id"]
    submitted, _ = service.record_generation(RecordGenerationRequest.model_validate({
        "project": PROJECT,
        "operation_id": f"{operation}-submitted",
        "change": {
            "kind": "update", "generation_record_id": generation_id,
            "expected_generation_revision": 1, "state": "submitted",
            "provider_ids": {"generation_ids": [f"offline-{operation}"]},
            "provider_response_metadata": {
                "format": "Deterministic local PCM fixture; no provider used.",
            },
        },
    }), owner_key=OWNER)
    completed, _ = service.record_generation(RecordGenerationRequest.model_validate({
        "project": PROJECT,
        "operation_id": f"{operation}-completed",
        "change": {
            "kind": "update", "generation_record_id": generation_id,
            "expected_generation_revision": submitted["generation"]["generation_revision"],
            "state": "completed",
            "provider_ids": {"generation_ids": [f"offline-{operation}"]},
        },
    }), owner_key=OWNER)
    pcm = synthetic_pcm(phase, frames)
    index = 1 if chapter_id == "chapter-one" else 2
    relative = f"Audiobook/Chapters/{index}/{operation}.pcm"
    source = service.root / relative
    source.parent.mkdir(parents=True, exist_ok=True)
    source.write_bytes(pcm)
    imported, replayed = service.import_audio(ImportAudioRequest.model_validate({
        "project": PROJECT,
        "operation_id": f"{operation}-import",
        "generation_record_id": generation_id,
        "expected_generation_revision": completed["generation"]["generation_revision"],
        "source": {"kind": "project_file", "filepath": relative,
                   "expected_sha256": hashlib.sha256(pcm).hexdigest()},
        "provenance": "native_generation",
        "source_format": {
            "container": "raw_pcm", "encoding": "signed_integer",
            "sample_rate_hz": 44_100, "channels": 1,
            "storage_bits": 16, "valid_bits": 16,
            "endianness": "little", "interleaving": "interleaved",
            "provider_format_evidence": "Deterministic local PCM fixture; no provider used.",
        },
    }), owner_key=OWNER)
    assert not replayed
    service.run_import_job(imported["job_id"])
    JOB_IDS.append(imported["job_id"])
    job = service.get_job(GetJobRequest(project=PROJECT, job_id=imported["job_id"]))
    assert job["state"] == "succeeded", job
    return job["result"]["take"]


def chapter_build(service: BookService, chapter_id: str, prepared: dict,
                  takes: list[dict], operation: str, head_revision: int | None) -> dict:
    state = ProjectState.discover(service.root)
    assert state is not None
    snapshot = state.snapshot(prepared["snapshot_id"])
    assert snapshot is not None
    build, replayed = service.build(BuildRequest.model_validate({
        "project": PROJECT,
        "operation_id": operation,
        "expected_head_revision": head_revision,
        "input": {
            "kind": "chapter", "chapter_id": chapter_id,
            "snapshot_id": prepared["snapshot_id"],
            "expected_manifest_revision": prepared["manifest_revision"],
            "request_plan_sha256": prepared["request_plan_sha256"],
            "takes": [{"chunk_id": take["chunk_id"], "take_id": take["take_id"],
                       "request_sha256": take["request_sha256"]} for take in takes],
        },
        "mode": "production_pcm",
        "outputs": {"master": True, "mp3_bitrate_kbps": 192},
        "gaps": [],
        "metadata": {"title": f"Synthetic {chapter_id}", "author": "Synthetic Example",
                     "edition": "16.0.0", "chapter_number": 1 if chapter_id == "chapter-one" else 2},
    }), owner_key=OWNER)
    assert not replayed
    service.run_build_job(build["job_id"])
    JOB_IDS.append(build["job_id"])
    job = service.get_job(GetJobRequest(project=PROJECT, job_id=build["job_id"]))
    assert job["state"] == "succeeded", job
    result = job["result"]
    build_path = (service.root / result["timeline_filepath"]).parent
    build_record = json.loads((build_path / "build.json").read_text(encoding="utf-8"))
    recipe = build_record["recipe"]
    assert canonical_json_sha256(recipe) == result["recipe_sha256"]
    assert recipe["mode"] == "production_pcm_mp3"
    assert recipe["tools"]["ffmpeg"]["executable"] == str(CognitaConfig().ffmpeg_executable.resolve())
    assert recipe["tools"]["ffprobe"]["executable"] == str(CognitaConfig().ffprobe_executable.resolve())
    assert recipe["outputs"] and all(item.get("bytes_sha256") for item in recipe["outputs"])
    assert {item["kind"] for item in result["outputs"]} == {"pcm_master", "mp3_download"}
    commit, replayed = service.commit_build(CommitBuildRequest.model_validate({
        "project": PROJECT,
        "operation_id": f"{operation}-commit",
        "build_id": result["build_id"],
        "expected_head_revision": head_revision,
        "intent": "accept_candidate",
        "acceptance": {
            "actor": "synthetic-example",
            "accepted_at": datetime.now(timezone.utc).isoformat(),
            "listening_review": "explicitly_waived",
            "notes": ["Synthetic fixture; no human listening or provider output."],
        },
    }), owner_key=OWNER)
    assert not replayed
    return {"build": result, "commit": commit, "recipe": recipe}


def book_build(service: BookService, operation: str,
               head_revision: int | None) -> dict:
    book = service.get_book(GetBookRequest.model_validate({
        "project": PROJECT, "book_id": "synthetic-book-16",
    }))
    assert not book["chapters_not_ready"], book["chapters_not_ready"]
    dependencies = book["current_chapter_dependencies"]
    candidate, replayed = service.build(BuildRequest.model_validate({
        "project": PROJECT,
        "operation_id": operation,
        "expected_head_revision": head_revision,
        "input": {
            "kind": "book", "book_id": "synthetic-book-16",
            "expected_layout_revision": book["layout_revision"],
            "chapters": dependencies,
        },
        "mode": "production_pcm",
        "outputs": {"master": True, "mp3_bitrate_kbps": 192},
        "gaps": [],
        "metadata": {"title": "Synthetic Offline Book", "author": "Synthetic Example",
                     "edition": "16.0.0"},
    }), owner_key=OWNER)
    assert not replayed
    service.run_build_job(candidate["job_id"])
    JOB_IDS.append(candidate["job_id"])
    finished = service.get_job(GetJobRequest(project=PROJECT, job_id=candidate["job_id"]))
    assert finished["state"] == "succeeded", finished
    result = finished["result"]
    timeline_path = service.root / result["timeline_filepath"]
    timeline = json.loads(timeline_path.read_text(encoding="utf-8"))
    recipe = json.loads((timeline_path.parent / "build.json").read_text(encoding="utf-8"))["recipe"]
    assert canonical_json_sha256(recipe) == result["recipe_sha256"]
    assert recipe["scope"] == "book" and recipe["mode"] == "production_pcm_mp3"
    assert recipe["tools"]["ffmpeg"]["executable"] == str(CognitaConfig().ffmpeg_executable.resolve())
    assert recipe["tools"]["ffprobe"]["executable"] == str(CognitaConfig().ffprobe_executable.resolve())
    state = ProjectState.discover(service.root)
    assert state is not None
    expected_inputs = []
    for dependency in dependencies:
        chapter_build = state.build(dependency["chapter_build_id"])
        assert chapter_build is not None
        pcm_output = next(item for item in chapter_build["result"]["outputs"]
                          if item["kind"] == "pcm_master")
        expected_inputs.append({
            "dependency": dependency, "filepath": pcm_output["filepath"],
            "bytes_sha256": pcm_output["bytes_sha256"], "media": pcm_output["media"],
        })
    assert recipe["chapter_inputs"] == expected_inputs
    assert {item["kind"] for item in result["outputs"]} == {"pcm_master", "mp3_download"}
    committed, replayed = service.commit_build(CommitBuildRequest.model_validate({
        "project": PROJECT,
        "operation_id": f"{operation}-commit",
        "build_id": result["build_id"],
        "expected_head_revision": head_revision,
        "intent": "accept_candidate",
        "acceptance": {
            "actor": "synthetic-example",
            "accepted_at": datetime.now(timezone.utc).isoformat(),
            "listening_review": "explicitly_waived",
            "notes": ["All media is deterministic synthetic fixture audio."],
        },
    }), owner_key=OWNER)
    assert not replayed
    return {"build": result, "commit": committed, "recipe": recipe, "timeline": timeline,
            "dependencies": dependencies}


def main() -> None:
    config = CognitaConfig()
    print(f"Cognita package: {cognita.__file__}")
    if config.ffmpeg_executable is None or config.ffprobe_executable is None:
        raise SystemExit(
            "This walkthrough needs the explicitly registered FFmpeg and ffprobe "
            "executables; the default container configuration did not find them."
        )
    with tempfile.TemporaryDirectory(prefix="cognita-book-16-example-") as temporary:
        root = Path(temporary) / "project"
        root.mkdir()
        prose_by_chapter = fixture_project(root)
        service = BookService(
            root, PROJECT,
            ffmpeg_executable=config.ffmpeg_executable,
            ffprobe_executable=config.ffprobe_executable,
        )

        takes_by_chapter = {}
        prepared_by_chapter = {}
        for index, chapter_id in enumerate(("chapter-one", "chapter-two")):
            prepared = prepare_chapter(service, chapter_id, prose_by_chapter[chapter_id],
                                       f"prepare-{chapter_id}-v1", "softly")
            prepared_by_chapter[chapter_id] = prepared
            takes_by_chapter[chapter_id] = [
                make_take(service, prepared, chapter_id, chunk["chunk_id"],
                          f"{chapter_id}-{chunk['chunk_id']}-v1", index * 10 + part)
                for part, chunk in enumerate(prepared["chunks"], start=1)
            ]

        first_acceptances = {}
        for chapter_id in ("chapter-one", "chapter-two"):
            first_acceptances[chapter_id] = chapter_build(
                service, chapter_id, prepared_by_chapter[chapter_id],
                takes_by_chapter[chapter_id], f"build-{chapter_id}-v1", None,
            )

        initial_book = book_build(service, "build-whole-book-v1", None)
        initial_book_two_start = next(
            int(entry["start_frame"]) for entry in initial_book["timeline"]["entries"]
            if entry["source_id"] == "chapter-two"
        )
        assert initial_book_two_start == 88_200
        quote = service.find_chunk(FindChunkRequest.model_validate({
            "project": PROJECT, "chapter_id": "chapter-one",
            "query": {"kind": "quote", "text": "lantern stayed lit"},
        }))
        assert quote["matches"][0]["chunk_ids"] == ["chapter-one-closing"]
        historic_time = service.find_chunk(FindChunkRequest.model_validate({
            "project": PROJECT, "chapter_id": "chapter-one",
            "query": {"kind": "timestamp", "build_id": first_acceptances["chapter-one"]["build"]["build_id"],
                      "seconds": 1.2},
        }))
        assert historic_time["matches"][0]["matched_build_id"] == first_acceptances["chapter-one"]["build"]["build_id"]
        assert historic_time["matches"][0]["chunk_ids"] == ["chapter-one-closing"]

        # Change one spoken-tag word without shifting the following chunk's
        # code-point range. The second chunk's exact take remains reusable.
        refreshed = prepare_chapter(
            service, "chapter-one", prose_by_chapter["chapter-one"],
            "prepare-chapter-one-retake", "gently",
            expected_manifest_revision=prepared_by_chapter["chapter-one"]["manifest_revision"],
        )
        unchanged = refreshed["chunks"][1]
        prior_second_take = takes_by_chapter["chapter-one"][1]
        assert unchanged["reuse_status"] == "reusable"
        assert prior_second_take["take_id"] in unchanged["reusable_take_ids"]
        prior_head = service.get_chapter(GetChapterRequest.model_validate({
            "project": PROJECT, "chapter_id": "chapter-one",
            "scope": {"kind": "production"},
        }))
        assert prior_head["current_outputs_stale"] is True
        assert prior_head["accepted_build_id"] == first_acceptances["chapter-one"]["build"]["build_id"]
        replacement = make_take(
            service, refreshed, "chapter-one", refreshed["chunks"][0]["chunk_id"],
            "chapter-one-retake-opening", 91, frames=88_200,
        )
        retake_takes = [replacement, prior_second_take]
        retake_acceptance = chapter_build(
            service, "chapter-one", refreshed, retake_takes,
            "build-chapter-one-retake", first_acceptances["chapter-one"]["commit"]["head_revision"],
        )
        assert retake_acceptance["commit"]["head_revision"] == 2
        previous_build = first_acceptances["chapter-one"]["build"]["build_id"]
        old_book = service.get_book(GetBookRequest.model_validate({
            "project": PROJECT, "book_id": "synthetic-book-16",
        }))
        assert old_book["accepted_build_id"] == initial_book["build"]["build_id"]
        assert old_book["current_outputs_stale"] is True
        retake_time = service.find_chunk(FindChunkRequest.model_validate({
            "project": PROJECT, "chapter_id": "chapter-one",
            "query": {"kind": "timestamp", "build_id": retake_acceptance["build"]["build_id"],
                      "seconds": 1.2},
        }))
        assert retake_time["matches"][0]["chunk_ids"] == ["chapter-one-opening"]
        assert retake_time["matches"][0]["matched_build_id"] == retake_acceptance["build"]["build_id"]

        second_book = book_build(
            service, "build-whole-book-v2", initial_book["commit"]["head_revision"],
        )
        second_book_two_start = next(
            int(entry["start_frame"]) for entry in second_book["timeline"]["entries"]
            if entry["source_id"] == "chapter-two"
        )
        assert second_book_two_start == 132_300
        assert second_book["dependencies"][1]["chapter_build_id"] == initial_book["dependencies"][1]["chapter_build_id"]

        rolled_back, _ = service.commit_build(CommitBuildRequest.model_validate({
            "project": PROJECT,
            "operation_id": "rollback-chapter-one-to-v1",
            "build_id": previous_build,
            "expected_head_revision": 2,
            "intent": "rollback",
            "acceptance": {
                "actor": "synthetic-example",
                "accepted_at": datetime.now(timezone.utc).isoformat(),
                "listening_review": "explicitly_waived",
                "notes": ["Restore the previously accepted synthetic chapter candidate."],
            },
        }), owner_key=OWNER)
        stale_book = service.get_book(GetBookRequest.model_validate({
            "project": PROJECT, "book_id": "synthetic-book-16",
        }))
        assert stale_book["accepted_build_id"] == second_book["build"]["build_id"]
        assert stale_book["current_outputs_stale"] is True
        restored_book_head, _ = service.commit_build(CommitBuildRequest.model_validate({
            "project": PROJECT,
            "operation_id": "rollback-book-to-v1",
            "build_id": initial_book["build"]["build_id"],
            "expected_head_revision": second_book["commit"]["head_revision"],
            "intent": "rollback",
            "acceptance": {
                "actor": "synthetic-example",
                "accepted_at": datetime.now(timezone.utc).isoformat(),
                "listening_review": "explicitly_waived",
                "notes": ["Restore the previously accepted synthetic book candidate."],
            },
        }), owner_key=OWNER)
        chapter = service.get_chapter(GetChapterRequest.model_validate({
            "project": PROJECT, "chapter_id": "chapter-one",
            "scope": {"kind": "production"}, "include_text": True,
        }))
        assert chapter["accepted_build_id"] == previous_build
        assert rolled_back["head_revision"] == 3
        assert restored_book_head["accepted_build_id"] == initial_book["build"]["build_id"]
        historic_match = service.find_chunk(FindChunkRequest.model_validate({
            "project": PROJECT,
            "chapter_id": "chapter-one",
            "query": {"kind": "quote", "text": "lantern stayed lit"},
        }))
        assert historic_match["matches"][0]["chunk_ids"] == ["chapter-one-closing"]

        generations = service.get_generations(GetGenerationsRequest.model_validate({
            "project": PROJECT,
            "query": {"kind": "chapter", "chapter_id": "chapter-one"},
            "limit": 100,
        }))
        generation_ids = [item["generation_record_id"] for item in generations["generations"]]
        assert len(generation_ids) == 3 and len(set(generation_ids)) == 3
        second_generations = service.get_generations(GetGenerationsRequest.model_validate({
            "project": PROJECT,
            "query": {"kind": "chapter", "chapter_id": "chapter-two"},
            "limit": 100,
        }))
        second_generation_ids = [item["generation_record_id"] for item in second_generations["generations"]]
        assert len(second_generation_ids) == 2 and len(set(second_generation_ids)) == 2
        for job_id in JOB_IDS:
            assert service.get_job(GetJobRequest(project=PROJECT, job_id=job_id))["state"] == "succeeded"

        # Snapshot only after all accepted jobs finish. The fixture helper backs
        # up the SQLite authority plus project files/media and the exact
        # external per-file de-index policy; it rejects nonquiesced state.
        del service
        gc.collect()
        data_dir = Path(temporary) / "project-data"
        data_dir.mkdir()
        source_master = "Project Files/Source/approved-source.docx"
        deindexed = DeindexedPaths(data_dir / "deindexed.json")
        assert deindexed.add(source_master)
        backup_root = Path(temporary) / "scoped-book-backup"
        snapshot = snapshot_quiesced_book_project(root, data_dir, backup_root)
        manifest_paths = {entry["path"] for entry in snapshot.entries}
        assert "data/deindexed.json" in manifest_paths
        assert "project/.cognita-storage/state.sqlite" in manifest_paths
        assert any(path.startswith("project/Audiobook/Chapters/1/builds/")
                   for path in manifest_paths)
        restored_root = Path(temporary) / "restored-project"
        restored_data = Path(temporary) / "restored-data"
        restore_quiesced_book_project(backup_root, restored_root, restored_data)
        restored_service = BookService(
            restored_root, PROJECT,
            ffmpeg_executable=config.ffmpeg_executable,
            ffprobe_executable=config.ffprobe_executable,
        )
        restored_book = restored_service.get_book(GetBookRequest.model_validate({
            "project": PROJECT, "book_id": "synthetic-book-16",
        }))
        assert restored_book["accepted_build_id"] == initial_book["build"]["build_id"]
        restored_generations = restored_service.get_generations(GetGenerationsRequest.model_validate({
            "project": PROJECT, "query": {"kind": "chapter", "chapter_id": "chapter-one"},
            "limit": 100,
        }))
        assert {item["generation_record_id"] for item in restored_generations["generations"]} == set(generation_ids)
        for job_id in JOB_IDS:
            assert restored_service.get_job(GetJobRequest(project=PROJECT, job_id=job_id))["state"] == "succeeded"
        restored_layout = restored_service.config().layout
        assert restored_layout is not None
        assert restored_layout.source_master_filepath == source_master
        assert restored_layout.index_policy.default_unknown == "exclude"
        state = ProjectState.discover(restored_root)
        assert state is not None
        assert not state.pending_publications("directory_move")
        restored_config = load_book_config(restored_root, state)
        assert restored_config.config_state == "enabled" and restored_config.layout is not None
        folder = state.folder_policy()
        restored_deindexed = DeindexedPaths(restored_data / "deindexed.json")
        assert restored_deindexed.sorted() == [source_master]
        assert restored_deindexed.load_error is None
        index_policy = EffectiveIndexPolicy(
            [FolderRule(path=path, indexed=indexed) for path, indexed in folder.rules],
            hard_exclusion_roots=(".cognita-storage",),
            deindexed_paths=restored_deindexed.sorted(),
            book_layout=restored_config.layout,
        )
        mutation_policy = BookMutationPolicy(
            restored_config.layout, config_state=restored_config.config_state,
            binding=restored_config.binding,
        )
        mutation = mutation_policy.decide(source_master, operation="write")
        assert not mutation.allowed and mutation.reason == "registered_source_master"
        source_bytes = (restored_root / source_master).read_bytes()
        source_sha256 = hashlib.sha256(source_bytes).hexdigest()
        source_read = restored_service.read_file(source_master, expected_bytes_sha256=source_sha256)
        assert source_read["bytes_sha256"] == source_sha256
        assert base64.b64decode(source_read["content_base64"]) == source_bytes
        source_listing = restored_service.list_files(
            "Project Files/Source", recursive=True,
            effective_index=index_policy,
            effective_read_only=lambda path: not mutation_policy.decide(path, operation="write").allowed,
        )
        listed_source = next(item for item in source_listing["entries"] if item["path"] == source_master)
        assert listed_source["effective_indexed"] is False
        assert listed_source["exclusion_reason"] == f"hard_exclusion:{source_master.casefold()}"
        assert listed_source["effective_read_only"] is True
        assert not index_policy.decision(source_master).indexed
        assert "project/Project Files/Source/approved-source.docx" in manifest_paths
        assert any(path.startswith("project/Audiobook/Chapters/1/builds/")
                   for path in manifest_paths)
        assert len(state.generations(chapter_id="chapter-one")) == 3
        assert len(state.generations(chapter_id="chapter-two")) == 2
        accepted_book_record = state.build(restored_book["accepted_build_id"])
        assert accepted_book_record is not None
        accepted_timeline = accepted_book_record["result"]["timeline_filepath"]
        accepted_recipe_path = restored_root / accepted_timeline
        restored_recipe = json.loads(
            (accepted_recipe_path.parent / "build.json").read_text(encoding="utf-8")
        )["recipe"]
        assert canonical_json_sha256(restored_recipe) == accepted_book_record["result"]["recipe_sha256"]
        manifest_hashes = {str(item["path"]): str(item["sha256"]) for item in snapshot.entries}
        for output in restored_book["exports"]:
            archive_path = f"project/{output['filepath']}"
            restored_path = restored_root / output["filepath"]
            actual_hash = hashlib.sha256(restored_path.read_bytes()).hexdigest()
            assert actual_hash == output["bytes_sha256"] == manifest_hashes[archive_path]
        print("Cognita 16 offline workflow passed: two chapters, native PCM imports, "
              "192 kbps chapter/book builds, exact recipes and changed timeline offsets, "
              "one-tag retake/reuse, stale history, quote/timestamp readers, chapter/book "
              "rollback, receipts, and verified scoped backup/restore.")
        del restored_service, state
        gc.collect()


if __name__ == "__main__":
    main()
