"""Document parsing for the 4.0 retrieval core (DESIGN-4.0-vector-engine.md D4.3).

Thin wrappers over the same libraries the 3.x engine used (pypdfium2 since 14.0, pymupdf before it,
python-docx, openpyxl, python-pptx), ported from mcp_server/ingestion.py with
the same format coverage and text-extraction shapes (PDF "[Page N]" markers,
DOCX headings as markdown, XLSX "## Sheet:" tables, ipynb code fences), so the
indexed text — and therefore retrieval quality — carries over.

Simplifications vs 3.x (deliberate):
- doc_id is content-addressed: sha256(source ":" content_sha256)[:16]. The 3.x
  id hashed path+mtime+size, so it churned on every touch; ours is stable
  across reindexes of unchanged content (indexes stay rebuildable, D4.10).
- Keyword extraction is config-route-based only (the CVE/IP/security-tool
  hardcoding served the upstream author's pentest corpus, not ours).
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import os
import re
import threading
from dataclasses import dataclass, field
from datetime import datetime, timezone
from fnmatch import fnmatch
from pathlib import Path
from typing import Callable

from .chunking import DEFAULT_CHUNK_OVERLAP, DEFAULT_CHUNK_SIZE, TextChunk, chunk_markdown, chunk_text
from .byte_facts import classify_text_bytes

log = logging.getLogger("cognita.parsing")

CODE_LANGUAGES = {
    ".py": "python",
    ".c": "c",
    ".h": "c",
    ".cpp": "cpp",
    ".js": "javascript",
    ".jsx": "javascript",
    ".ts": "typescript",
    ".tsx": "typescript",
}

# ---------------------------------------------------------------------------
# Tiers (4.4). A file is either EMBEDDED (chunked + vectorized + hybrid-
# searchable — everything 4.3 did) or REGISTERED (stored whole, keyword-
# searchable, never embedded). See DESIGN-4.4-registered-tier.md.
#
# Registered exists because code is the worst possible embedding input: the
# splitter cuts mid-function and the resulting vectors are mush that pollutes
# every prose query. Exact-token matching on the same text is precise and
# noise-free, so the fix is to drop the semantic half only, not the file.
# ---------------------------------------------------------------------------

TIER_EMBEDDED = "embedded"
TIER_REGISTERED = "registered"

# Prose + structured documents: these chunk and embed well.
DEFAULT_INDEXED_EXTENSIONS = frozenset({
    ".md", ".txt", ".pdf", ".json", ".xml", ".docx", ".xlsx", ".pptx", ".csv", ".ipynb",
})

# Code and config. .py/.js/.ts/... were EMBEDDED through 4.3 and are MOVED here
# by 4.4 — this is the point of the tier, not an addition beside it. The first
# reindex after upgrade purges their existing chunks and vectors.
#
# .json is deliberately NOT here: structured prose and application exports
# can be content meant for semantic search. It stays embedded.
DEFAULT_REGISTERED_EXTENSIONS = frozenset({
    ".py", ".sh", ".js", ".css", ".yml", ".yaml", ".toml", ".ini", ".sql", ".json5",
    *CODE_LANGUAGES,  # .c .h .cpp .jsx .ts .tsx — same argument as .py
})

# Backwards-compatible union: every extension Cognita will touch by default.
SUPPORTED_FORMATS = DEFAULT_INDEXED_EXTENSIONS | DEFAULT_REGISTERED_EXTENSIONS


def normalize_extensions(extensions) -> frozenset[str]:
    """Lowercase, dot-prefix and de-blank a configured extension list."""
    out = set()
    for raw in extensions or ():
        ext = str(raw).strip().lower()
        if not ext:
            continue
        out.add(ext if ext.startswith(".") else f".{ext}")
    return frozenset(out)


@dataclass(frozen=True, slots=True)
class ExtensionPolicy:
    """Which extensions land in which tier, for ONE project.

    Per-project by construction (D4.4-2): Doug's connectors have very different
    contents and a scripts-heavy one must not make a prose one start hoovering
    up source files.
    """

    embedded: frozenset[str] = DEFAULT_INDEXED_EXTENSIONS
    registered: frozenset[str] = DEFAULT_REGISTERED_EXTENSIONS
    conflicts: tuple[str, ...] = ()  # in BOTH lists; registered won

    @classmethod
    def build(cls, indexed=None, registered=None) -> "ExtensionPolicy":
        """Resolve two configured lists into a policy. An extension named in
        both tiers resolves to REGISTERED and is reported in .conflicts so the
        caller can warn at startup.

        Registered wins deliberately. The extensions this tier exists for are
        ALREADY in the embedded allowlist, so "embedded wins" would make adding
        one to registered_extensions do nothing at all — the feature would ship
        and be a no-op. Precedence is the safety net; the mechanism is that the
        defaults above move code out of the embedded list outright.
        """
        emb = (DEFAULT_INDEXED_EXTENSIONS if indexed is None
               else normalize_extensions(indexed))
        reg = (DEFAULT_REGISTERED_EXTENSIONS if registered is None
               else normalize_extensions(registered))
        overlap = emb & reg
        return cls(embedded=emb - overlap, registered=reg,
                   conflicts=tuple(sorted(overlap)))

    def tier_for(self, suffix: str) -> str | None:
        """The tier a file suffix belongs to, or None if it is not indexed."""
        ext = suffix.lower()
        if ext in self.embedded:
            return TIER_EMBEDDED
        if ext in self.registered:
            return TIER_REGISTERED
        return None

    @property
    def all_extensions(self) -> frozenset[str]:
        return self.embedded | self.registered


DEFAULT_POLICY = ExtensionPolicy()


@dataclass(slots=True)
class ParsedDocument:
    """A parsed file, ready for chunking + embedding + storage."""

    source: str  # path relative to documents_dir, forward slashes
    format: str  # file suffix, e.g. ".md"
    content: str
    content_hash: str  # sha256 hex of the extracted text
    doc_id: str  # sha256(source:content_hash)[:16] — stable while content is
    category: str = "general"
    keywords: list[str] = field(default_factory=list)
    file_mtime: datetime | None = None
    file_size: int | None = None
    tier: str = TIER_EMBEDDED
    # In-flight source evidence for the retrieval core.  It is never stored in
    # the search database; book indexing uses it to bind parsing to one read.
    captured_raw: bytes | None = field(default=None, repr=False, compare=False)
    book_index_context: object | None = field(default=None, repr=False, compare=False)
    book_index_record: object | None = field(default=None, repr=False, compare=False)

    @property
    def is_registered(self) -> bool:
        return self.tier == TIER_REGISTERED

    def chunks(
        self,
        chunk_size: int = DEFAULT_CHUNK_SIZE,
        chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
    ) -> list[TextChunk]:
        """Chunks for the embedded tier. Registered documents are stored whole
        and never chunked — returning [] here is what makes the reindex cost for
        that tier filesystem-only (no splitter pass, no embedding call)."""
        if self.is_registered:
            return []
        if self.format == ".md":
            return chunk_markdown(self.content, chunk_size, chunk_overlap)
        return chunk_text(self.content, chunk_size, chunk_overlap)


def compute_doc_id(source: str, content_hash: str) -> str:
    return hashlib.sha256(f"{source}:{content_hash}".encode()).hexdigest()[:16]


# ---------------------------------------------------------------------------
# Format extractors.  parse_file uses byte-backed bodies below; path wrappers
# remain narrow compatibility helpers for direct diagnostics/tests.
# ---------------------------------------------------------------------------


def _read_text_bytes(raw: bytes) -> str:
    """Decode a text file, sniffing the BOM before assuming UTF-8.

    This was `read_text(encoding="utf-8", errors="ignore")` with no sniff, which
    fails in two directions on real files:

      * A UTF-16 file — Excel's "Unicode Text" export, a Windows editor's default
        — decodes to text with a NUL between every character. PostgreSQL `text`
        cannot hold U+0000, so replace_document failed inside its transaction and
        the document was recorded in summary["errors"] as an opaque database
        error. The file was simply never indexed and the reason was unreadable to
        anyone who had not written this code.
      * errors="ignore" on latin-1 content silently DROPS bytes, so content_hash
        described a mangled string and search saw text the file does not contain.
        That is quieter and worse.

    UTF-8 is still the default and the overwhelmingly common case; the BOM sniff
    only redirects files that say outright what they are. NULs are stripped
    defensively afterwards so a mis-sniffed file degrades to bad text rather than
    to a failed transaction.
    """
    # Keep the historical BOM behavior for this low-level compatibility helper
    # (some callers use it to inspect UTF-16 exports), while all UTF-8 indexing
    # uses the shared 9.2 tolerant view below.
    if not raw.startswith((b"\xff\xfe", b"\xfe\xff", b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        view = classify_text_bytes(raw)
        if not view.accepted:
            raise ValueError(f"{view.reason}: {view.message}")
        return view.indexed_text
    for bom, encoding in (
        (b"\xff\xfe\x00\x00", "utf-32"), (b"\x00\x00\xfe\xff", "utf-32"),
        (b"\xff\xfe", "utf-16"), (b"\xfe\xff", "utf-16"),
        (b"\xef\xbb\xbf", "utf-8-sig"),
    ):
        if raw.startswith(bom):
            text = raw.decode(encoding, errors="ignore")
            break
    else:
        text = raw.decode("utf-8", errors="ignore")
    if "\x00" in text:
        text = text.replace("\x00", "")
    if text.startswith("﻿"):
        text = text[1:]
    # Universal-newline translation, which Path.read_text did for free in text
    # mode and read_bytes does not. Dropping it would have let CRLF into the
    # INDEXED copy — breaking the frontmatter regex, the section splitter and
    # every anchor built from indexed text. Normalization belongs on the indexed
    # copy and never on the stored file, which _write_verbatim still writes byte
    # for byte.
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _read_text(filepath: Path) -> str:
    """Compatibility helper for callers that intentionally parse a path."""
    return _read_text_bytes(filepath.read_bytes())


def _validate_text_bytes(raw: bytes) -> None:
    """Reject incompatible/binary bytes before a text extractor can mutate them."""
    view = classify_text_bytes(raw)
    if not view.accepted:
        raise ValueError(f"{view.reason}: {view.message}")


def _extract_markdown_bytes(raw: bytes) -> str:
    content = _read_text_bytes(raw)
    # Strip YAML frontmatter, as 3.x did — metadata noise, not prose.
    frontmatter = re.match(r"^---\n(.*?)\n---\n", content, re.DOTALL)
    if frontmatter:
        content = content[frontmatter.end():]
    return content


def _extract_markdown(filepath: Path) -> str:
    return _extract_markdown_bytes(filepath.read_bytes())


# PDFium is not thread-safe (pypdfium2 README, "Incompatibility with Threading"), and
# Cognita parses from worker threads (asyncio.to_thread in the walk's parse-ahead
# producer, index_file, move_file and get_document), so two PDFs can be parsed at
# once. One call to _extract_pdf holds this lock from open to close; PDF extraction is
# therefore serialized across threads. Parsing is small next to embedding, so it costs
# little (DESIGN-14.0 §1.3).
_PDFIUM_LOCK = threading.Lock()


def _extract_pdf_bytes(data: bytes, *, name: str) -> str:
    """PDF text through pypdfium2 (DESIGN-14.0 §1). Shape unchanged from the PyMuPDF
    version: "[Page N]\\n<text>" per page that has text, N counting skipped pages,
    joined by a blank line.

    Failures keep PyMuPDF's classes so every caller answers as before: a user-password
    PDF raises ValueError("document closed or encrypted") (get_document maps ValueError
    to unsupported_format), a zero-page PDF returns "" (parse_file then returns None),
    and any other PdfiumError propagates as the RuntimeError it is.
    """
    import pypdfium2 as pdfium  # imported lazily; keeps the import off the startup path
    import pypdfium2.raw as pdfium_raw

    # `data` must stay alive until the document is closed: PDFium reads from
    # this buffer, it does not copy it.
    with _PDFIUM_LOCK:
        # Loaded through the raw API, not PdfDocument(path_or_stream). pypdfium2
        # 5.13 treats "failed to load" and "loaded, zero pages" alike: it raises
        # with FPDF_GetLastError(), and PDFium does NOT reset that code on a
        # successful load. Seen on kei 2026-09-28: a zero-page PDF parsed right
        # after a corrupt one reported the corrupt file's "data format error".
        # pypdfium2 also never closed the zero-page document it had loaded, so the
        # file handle (or native document) leaked on every parse. Here the last
        # error is read only when the load itself returned NULL, which is when
        # PDFium sets it, and a zero-page document is closed by us.
        raw_doc = pdfium_raw.FPDF_LoadMemDocument64(data, len(data), None)
        if not raw_doc:
            err_code = pdfium_raw.FPDF_GetLastError()
            exc = pdfium.PdfiumError(
                f"Failed to load document (PDFium error {err_code}).", err_code=err_code
            )
            log.warning("pdf.open.failed file=%s exc=%s err_code=%s",
                        name, type(exc).__name__, err_code)
            if err_code == pdfium_raw.FPDF_ERR_PASSWORD:
                raise ValueError("document closed or encrypted") from exc
            raise exc
        if pdfium_raw.FPDF_GetPageCount(raw_doc) < 1:
            # Loaded fine but has no pages: nothing to index, not an error.
            pdfium_raw.FPDF_CloseDocument(raw_doc)
            return ""
        doc = pdfium.PdfDocument(raw_doc)  # takes ownership; doc.close() frees it
        parts = []
        try:
            for page_index in range(len(doc)):
                page = None
                try:
                    page = doc.get_page(page_index)
                    textpage = page.get_textpage()
                    try:
                        # Bounded to the page box, as PyMuPDF's default clip was;
                        # get_text_range() can include off-page text.
                        text = textpage.get_text_bounded()
                    finally:
                        textpage.close()
                except pdfium.PdfiumError as exc:
                    log.warning("pdf.page.failed file=%s page=%d exc=%s err_code=%s",
                                name, page_index + 1, type(exc).__name__, exc.err_code)
                    raise
                finally:
                    if page is not None:
                        page.close()
                # PostgreSQL text cannot hold NUL (_read_text strips it too).
                text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
                if text.strip():
                    parts.append(f"[Page {page_index + 1}]\n{text}")
        finally:
            doc.close()
    return "\n\n".join(parts)


def _extract_pdf(filepath: Path) -> str:
    return _extract_pdf_bytes(filepath.read_bytes(), name=filepath.name)


def _extract_json_bytes(raw: bytes) -> str:
    text = _read_text_bytes(raw)
    try:
        return json.dumps(json.loads(text), indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        return text


def _extract_json(filepath: Path) -> str:
    return _extract_json_bytes(filepath.read_bytes())


def _extract_docx_bytes(raw: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(raw))
    parts = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style and para.style.name.startswith("Heading"):
            try:
                level = int(para.style.name.split()[-1])
                parts.append(f"{'#' * level} {text}")
            except (ValueError, IndexError):
                parts.append(f"## {text}")
        else:
            parts.append(text)
    for table in doc.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if rows:
            parts.append("\n".join(rows))
    return "\n\n".join(parts)


def _extract_docx(filepath: Path) -> str:
    return _extract_docx_bytes(filepath.read_bytes())


def _extract_xlsx_bytes(raw: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    parts = []
    try:
        for sheet_name in wb.sheetnames:
            parts.append(f"## Sheet: {sheet_name}")
            for row in wb[sheet_name].iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                line = " | ".join(cells).strip()
                if line and set(line) != {" ", "|"}:
                    parts.append(line)
    finally:
        wb.close()
    return "\n\n".join(parts)


def _extract_xlsx(filepath: Path) -> str:
    return _extract_xlsx_bytes(filepath.read_bytes())


def _extract_pptx_bytes(raw: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(raw))
    parts = []
    for i, slide in enumerate(prs.slides):
        texts = [
            para.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame
            for para in shape.text_frame.paragraphs
            if para.text.strip()
        ]
        if texts:
            parts.append(f"## Slide {i + 1}\n" + "\n".join(texts))
    return "\n\n".join(parts)


def _extract_pptx(filepath: Path) -> str:
    return _extract_pptx_bytes(filepath.read_bytes())


def _extract_csv_bytes(raw: bytes) -> str:
    rows = list(csv.reader(io.StringIO(_read_text_bytes(raw))))
    return "\n".join(" | ".join(row) for row in rows)


def _extract_csv(filepath: Path) -> str:
    return _extract_csv_bytes(filepath.read_bytes())


def _extract_ipynb_bytes(raw: bytes) -> str:
    text = _read_text_bytes(raw)
    try:
        nb = json.loads(text)
    except json.JSONDecodeError:
        return text
    parts = []
    for cell in nb.get("cells", []):
        source = cell.get("source", "")
        if isinstance(source, list):
            source = "".join(source)
        if not source.strip():
            continue
        if cell.get("cell_type") == "markdown":
            parts.append(source)
        elif cell.get("cell_type") == "code":
            parts.append(f"```python\n{source}\n```")
    return "\n\n".join(parts)


def _extract_ipynb(filepath: Path) -> str:
    return _extract_ipynb_bytes(filepath.read_bytes())


# parse_file owns a single captured source buffer, so index extraction never
# validates one read and extracts another.
_BYTE_EXTRACTORS = {
    ".md": _extract_markdown_bytes,
    ".txt": _read_text_bytes,
    ".pdf": _extract_pdf_bytes,
    ".json": _extract_json_bytes,
    ".xml": _read_text_bytes,
    ".docx": _extract_docx_bytes,
    ".xlsx": _extract_xlsx_bytes,
    ".pptx": _extract_pptx_bytes,
    ".csv": _extract_csv_bytes,
    ".ipynb": _extract_ipynb_bytes,
    **{ext: _read_text_bytes for ext in CODE_LANGUAGES},
}


# ---------------------------------------------------------------------------
# Category + keywords (config-driven; default: everything "general")
# ---------------------------------------------------------------------------


def detect_category(rel_path: str, category_mappings: dict[str, str]) -> str:
    """Longest-match substring lookup on the relative path, as in 3.x."""
    path_str = rel_path.replace("\\", "/").lower()
    for pattern, category in sorted(category_mappings.items(), key=lambda x: len(x[0]), reverse=True):
        if pattern in path_str:
            return category
    return "general"


def extract_keywords(content: str, keyword_routes: dict[str, list[str]]) -> list[str]:
    content_lower = content.lower()
    found = {
        kw.lower()
        for route_keywords in keyword_routes.values()
        for kw in route_keywords
        if kw.lower() in content_lower
    }
    return sorted(found)


# ---------------------------------------------------------------------------
# Entry points
# ---------------------------------------------------------------------------


def parse_file(
    filepath: Path,
    documents_dir: Path,
    *,
    category_mappings: dict[str, str] | None = None,
    keyword_routes: dict[str, list[str]] | None = None,
    policy: ExtensionPolicy | None = None,
    captured_content: Callable[[str, str, bytes], object] | None = None,
) -> ParsedDocument | None:
    """Parse one file. Returns None for empty documents (nothing to index).

    Raises ValueError for unsupported formats and lets parser-library errors
    propagate — the directory walker catches per-file failures so one bad file
    can't sink a project reindex.

    The policy decides the tier. Registered-tier files fall back to a plain
    text read when no dedicated extractor exists (.sh/.css/.toml/... are just
    text), which is exactly what a keyword-only tier needs.
    """
    policy = policy or DEFAULT_POLICY
    suffix = filepath.suffix.lower()
    tier = policy.tier_for(suffix)
    if tier is None:
        raise ValueError(f"Unsupported format: {suffix}")
    extractor = _BYTE_EXTRACTORS.get(suffix)
    if extractor is None:
        if tier != TIER_REGISTERED:
            raise ValueError(f"Unsupported format: {suffix}")
        extractor = _read_text_bytes
    # Stat BEFORE the read, never after. The stat persisted here is exactly what
    # the next smart reindex compares against (RetrievalCore._stat_matches), so it
    # must never describe a NEWER file than the content we actually extracted. A
    # write landing between the two calls would otherwise commit the old text under
    # the new file's mtime+size, and every later smart reindex would skip the file
    # as unchanged — a stale row recoverable only by a full_rebuild. Statting first
    # fails the safe way round: a stale stat costs one redundant reindex and heals.
    stat = filepath.stat()
    raw = filepath.read_bytes()
    try:
        source = filepath.relative_to(documents_dir).as_posix()
    except ValueError:
        source = filepath.as_posix()
    # Dedicated binary formats (PDF/Office) run through their parser first;
    # only extensions that promise ordinary text use the conservative byte
    # classifier.
    if suffix not in {".pdf", ".docx", ".xlsx", ".pptx"}:
        _validate_text_bytes(raw)
    selected = captured_content(source, suffix, raw) if captured_content is not None else None
    book_index_context = None
    if isinstance(selected, tuple):
        content, book_index_context = selected
    else:
        content = selected
    if content is None:
        content = _extract_pdf_bytes(raw, name=filepath.name) if suffix == ".pdf" else extractor(raw)
    if not content or not content.strip():
        return None
    content_hash = hashlib.sha256(content.encode()).hexdigest()
    return ParsedDocument(
        source=source,
        format=suffix,
        content=content,
        content_hash=content_hash,
        doc_id=compute_doc_id(source, content_hash),
        category=detect_category(source, category_mappings or {}),
        keywords=extract_keywords(content, keyword_routes or {}),
        file_mtime=datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc),
        file_size=stat.st_size,
        tier=tier,
        # The project callback has already bound any role/configuration facts;
        # ParsedDocument keeps extracted text and immutable facts, never media
        # or source buffers, while it waits for embedding/publication.
        captured_raw=None,
        book_index_context=book_index_context,
    )


# Cloud-sync conflict copies (5.0 §10). Cognita's documents root on kei is a
# bidirectionally synced OneDrive folder, so a second writer exists that nothing
# in the tool surface ever acknowledged. When that writer loses a race it does
# not fail — it writes a SECOND file beside the first, named after the device
# that lost. Indexed, that copy is a silent wrong answer: a builder glob over
# `*.txt` picks it up alongside the real file and the pack output changes with
# nobody having edited anything.
#
# These are fnmatch globs over the file NAME, config-overridable via
# `sync_conflict_patterns`. They are deliberately conservative — the cost of a
# false positive is a real document going unindexed — and every skip is logged
# and counted into get_index_stats().sync_conflicts, so an exclusion that should
# not have happened is discoverable instead of silent. There is no filename rule
# that is free of false positives; being loud is the mitigation.
SYNC_CONFLICT_PATTERNS: list[str] = [
    "*-*-conflict.*",  # OneDrive: notes-DESKTOP-conflict.txt
    "*-*-conflicted.*",
    "*conflicted copy*",  # OneDrive/Dropbox English: "note (X's conflicted copy 2026-08-29).md"
    "*.sync-conflict-*",  # Syncthing: note.sync-conflict-20260829-120000-ABCDEFG.md
    # abraunegg `onedrive` (the Linux client on kei): two.md -> two-kei-safeBackup-0001.md.
    # 14.0.1: missing until then, so kei's own conflict copies were indexed as documents;
    # three sat in the corpus and one broke the 14.0.0 web self-test's pack checks.
    "*-*-safebackup-[0-9]*",
]


def is_sync_conflict(name: str, patterns: list[str] | None = None) -> bool:
    """True when `name` looks like a cloud-sync conflict copy, not a document."""
    return any(
        fnmatch(name.lower(), pattern.lower())
        for pattern in (SYNC_CONFLICT_PATTERNS if patterns is None else patterns)
    )


def partition_sync_conflicts(
    files: list[Path], patterns: list[str] | None = None
) -> tuple[list[Path], list[Path]]:
    """Split a walk's output into (indexable, sync-conflict copies)."""
    kept: list[Path] = []
    conflicts: list[Path] = []
    for filepath in files:
        (conflicts if is_sync_conflict(filepath.name, patterns) else kept).append(filepath)
    return kept, conflicts


def _is_excluded(rel_path: Path, patterns: list[str]) -> bool:
    rel_str = rel_path.as_posix()
    for pattern in patterns:
        if fnmatch(rel_str, pattern) or any(fnmatch(part, pattern) for part in rel_path.parts):
            return True
    return False


def iter_document_files(
    documents_dir: Path,
    exclude_patterns: list[str],
    policy: ExtensionPolicy | None = None,
    *,
    raise_on_error: bool = False,
) -> list[Path]:
    """All indexable files under documents_dir, excludes applied, sorted.

    "Indexable" spans BOTH tiers — the walk collects registered-extension files
    too; parse_file/the retrieval core decide what happens to each.

    Follows symlinks with cycle protection (as 3.x did — OneDrive trees on kei
    contain them).
    """
    # os.walk swallows EVERY error by default, including "the root does not
    # exist", and yields nothing. That empty list used to flow into
    # index_project's removal sweep, which reads "no files on disk" as "every
    # document was deleted" and drops the whole project index. An unreachable
    # tree is an ERROR, not an empty tree — say so here, at the only place that
    # can still tell the two apart.
    if not documents_dir.is_dir():
        raise NotADirectoryError(
            f"documents_dir is not a readable directory: {documents_dir}"
        )

    def _on_walk_error(exc: OSError) -> None:
        if raise_on_error:
            raise exc
        # A subdirectory we cannot read must not silently shrink the file list
        # either; log it loudly and let index_project's guard catch the case
        # where the shortfall is total.
        log.warning("Skipping unreadable path during walk: %s", exc)

    known = (policy or DEFAULT_POLICY).all_extensions
    files: list[Path] = []
    seen_dirs: set[str] = set()
    for root, dirs, names in os.walk(documents_dir, followlinks=True, onerror=_on_walk_error):
        real_root = os.path.realpath(root)
        if real_root in seen_dirs:
            dirs.clear()
            continue
        seen_dirs.add(real_root)
        root_path = Path(root)
        dirs[:] = [
            d for d in dirs
            if not _is_excluded((root_path / d).relative_to(documents_dir), exclude_patterns)
        ]
        for name in names:
            filepath = root_path / name
            if filepath.suffix.lower() not in known:
                continue
            if _is_excluded(filepath.relative_to(documents_dir), exclude_patterns):
                continue
            files.append(filepath)
    return sorted(files)
