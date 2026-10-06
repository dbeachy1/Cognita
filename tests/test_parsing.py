"""Unit tests for cognita.parsing — format extraction, ids, discovery, excludes."""

import hashlib
import json
from pathlib import Path

import pytest

from cognita.parsing import (
    _read_text,
    DEFAULT_INDEXED_EXTENSIONS,
    DEFAULT_REGISTERED_EXTENSIONS,
    SUPPORTED_FORMATS,
    TIER_EMBEDDED,
    TIER_REGISTERED,
    ExtensionPolicy,
    compute_doc_id,
    detect_category,
    extract_keywords,
    iter_document_files,
    parse_file,
)


@pytest.fixture
def docs(tmp_path):
    d = tmp_path / "docs"
    d.mkdir()
    return d


def test_parse_markdown_strips_frontmatter(docs):
    f = docs / "note.md"
    f.write_text("---\ntitle: x\n---\n# Real Title\n\nbody text\n", encoding="utf-8")
    doc = parse_file(f, docs)
    assert doc.content.startswith("# Real Title")
    assert doc.format == ".md"
    assert doc.source == "note.md"
    assert doc.category == "general"
    assert doc.file_size == f.stat().st_size
    assert len(doc.doc_id) == 16


def test_parse_nested_source_is_posix_relative(docs):
    sub = docs / "Project Files"
    sub.mkdir()
    f = sub / "plan.md"
    f.write_text("# Plan\n\ncontent here", encoding="utf-8")
    doc = parse_file(f, docs)
    assert doc.source == "Project Files/plan.md"


def test_doc_id_is_content_addressed(docs):
    f = docs / "a.txt"
    f.write_text("same content", encoding="utf-8")
    doc1 = parse_file(f, docs)
    # Touch the file (mtime changes, content doesn't) — id must not churn (unlike 3.x)
    f.touch()
    doc2 = parse_file(f, docs)
    assert doc1.doc_id == doc2.doc_id
    f.write_text("different content", encoding="utf-8")
    assert parse_file(f, docs).doc_id != doc1.doc_id
    # Same content at a different path is a different document
    g = docs / "b.txt"
    g.write_text("same content", encoding="utf-8")
    assert parse_file(g, docs).doc_id != doc1.doc_id
    assert compute_doc_id("a.txt", "h") != compute_doc_id("b.txt", "h")


def test_stat_is_never_newer_than_the_content_it_describes(docs, monkeypatch):
    """A write landing mid-parse must not poison the row's mtime+size.

    The stat parse_file persists is what RetrievalCore._stat_matches compares
    against on the next smart reindex. If it described the post-write file while
    the chunks came from the pre-write text, every later smart reindex would skip
    the file as unchanged and the new content would stay invisible until someone
    ran a full_rebuild. Statting before the read makes the failure self-healing.
    """
    from cognita import parsing

    original = "# Original\n\nthe text we actually extracted\n"
    f = docs / "raced.md"
    f.write_text(original, encoding="utf-8")
    pre_size = f.stat().st_size  # not len(original): Windows writes \n as \r\n
    real = parsing._BYTE_EXTRACTORS[".md"]

    def racing_extractor(raw):
        text = real(raw)  # parse_file has already captured the source bytes...
        f.write_text(original + "\nappended by a concurrent writer\n", encoding="utf-8")
        return text  # ...and a write lands before we return

    monkeypatch.setitem(parsing._BYTE_EXTRACTORS, ".md", racing_extractor)
    doc = parse_file(f, docs)

    assert doc.content.startswith("# Original")
    assert "concurrent writer" not in doc.content  # we did not read the new bytes
    # ...so the stat must not claim we did: it describes the file as it was read,
    # which no longer matches disk — that mismatch is what re-reads it next pass.
    assert doc.file_size == pre_size
    assert doc.file_size != f.stat().st_size


def test_parse_file_stats_before_one_captured_read(docs, monkeypatch):
    """Validation and extraction must consume one buffer, not reopen a changed path."""
    f = docs / "captured.md"
    original = "# Captured\n\nfirst bytes only\n"
    replacement = "# Replacement\n\nnew bytes must not be indexed\n"
    f.write_text(original, encoding="utf-8")
    before = f.stat()
    path_type = type(f)
    original_read_bytes = path_type.read_bytes
    reads: list[Path] = []

    def replace_after_capture(path):
        if path == f:
            reads.append(path)
            raw = original_read_bytes(path)
            f.write_text(replacement, encoding="utf-8")
            return raw
        return original_read_bytes(path)

    monkeypatch.setattr(path_type, "read_bytes", replace_after_capture)
    parsed = parse_file(f, docs)

    assert reads == [f]
    assert parsed.captured_raw is None
    assert parsed.content == original
    assert parsed.content_hash == hashlib.sha256(original.encode()).hexdigest()
    assert parsed.file_size == before.st_size
    assert parsed.file_size != f.stat().st_size


def test_parse_file_docx_matches_existing_office_extractor(docs):
    """The BytesIO path preserves the established DOCX heading/table shape."""
    import docx
    from cognita.parsing import _extract_docx

    path = docs / "office.docx"
    document = docx.Document()
    document.add_heading("A heading", level=2)
    document.add_paragraph("ordinary paragraph")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "left"
    table.rows[0].cells[1].text = "right"
    document.save(path)

    assert parse_file(path, docs).content == _extract_docx(path)


def test_parse_file_xlsx_and_pptx_match_existing_office_extractors(docs):
    """All Office loaders accept the captured in-memory source stream."""
    import openpyxl
    from pptx import Presentation
    from cognita.parsing import _extract_pptx, _extract_xlsx

    xlsx_path = docs / "sheet.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Data"
    workbook.active.append(["name", "count"])
    workbook.active.append(["bolt", 4])
    workbook.save(xlsx_path)
    workbook.close()

    pptx_path = docs / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Title"
    slide.placeholders[1].text = "body"
    presentation.save(pptx_path)

    assert parse_file(xlsx_path, docs).content == _extract_xlsx(xlsx_path)
    assert parse_file(pptx_path, docs).content == _extract_pptx(pptx_path)


def test_empty_file_parses_to_none(docs):
    f = docs / "empty.md"
    f.write_text("   \n", encoding="utf-8")
    assert parse_file(f, docs) is None


def test_unsupported_format_raises(docs):
    f = docs / "archive.zip"
    f.write_bytes(b"PK\x03\x04")
    with pytest.raises(ValueError, match="Unsupported format"):
        parse_file(f, docs)


def test_parse_json_pretty_prints(docs):
    f = docs / "data.json"
    f.write_text('{"b":1,"a":[2,3]}', encoding="utf-8")
    doc = parse_file(f, docs)
    assert json.loads(doc.content) == {"b": 1, "a": [2, 3]}
    assert "\n" in doc.content  # indented


def test_parse_code_reads_verbatim(docs):
    f = docs / "tool.py"
    f.write_text("def main():\n    pass\n", encoding="utf-8")
    doc = parse_file(f, docs)
    assert doc.content == "def main():\n    pass\n"
    assert doc.format == ".py"


def test_parse_csv_renders_rows(docs):
    f = docs / "table.csv"
    f.write_text("name,qty\nbolt,4\n", encoding="utf-8")
    doc = parse_file(f, docs)
    assert doc.content == "name | qty\nbolt | 4"


def test_parse_ipynb_extracts_cells(docs):
    nb = {
        "cells": [
            {"cell_type": "markdown", "source": ["# Notes\n"]},
            {"cell_type": "code", "source": "print('hi')"},
            {"cell_type": "code", "source": "", "outputs": ["ignored"]},
        ]
    }
    f = docs / "nb.ipynb"
    f.write_text(json.dumps(nb), encoding="utf-8")
    doc = parse_file(f, docs)
    assert "# Notes" in doc.content
    assert "```python\nprint('hi')\n```" in doc.content


def test_iter_document_files_applies_excludes_and_formats(docs):
    (docs / "a.md").write_text("x", encoding="utf-8")
    (docs / "b.zip").write_bytes(b"PK")  # unsupported → skipped
    (docs / "run.sh").write_text("#!/bin/sh", encoding="utf-8")  # registered tier (4.4)
    backups = docs / "backups"
    backups.mkdir()
    (backups / "old.md").write_text("x", encoding="utf-8")  # excluded dir
    nested = docs / "sub" / "backups"
    nested.mkdir(parents=True)
    (nested / "deep.md").write_text("x", encoding="utf-8")  # excluded at any depth
    (docs / "sub" / "keep.md").write_text("x", encoding="utf-8")

    # The walk spans BOTH tiers (4.4): run.sh is collected as a registered
    # document; only genuinely unknown extensions (.zip) are skipped.
    files = iter_document_files(docs, ["backups"])
    rel = [f.relative_to(docs).as_posix() for f in files]
    assert rel == ["a.md", "run.sh", "sub/keep.md"]


def test_iter_document_files_honors_per_project_policy(docs):
    """Per-connector scoping: a project that does not register .sh must not
    collect it, even though .sh is in the global default."""
    (docs / "a.md").write_text("x", encoding="utf-8")
    (docs / "run.sh").write_text("#!/bin/sh", encoding="utf-8")

    prose_only = ExtensionPolicy.build([".md"], [])
    rel = [f.relative_to(docs).as_posix()
           for f in iter_document_files(docs, ["backups"], prose_only)]
    assert rel == ["a.md"]

    with_scripts = ExtensionPolicy.build([".md"], [".sh"])
    rel = [f.relative_to(docs).as_posix()
           for f in iter_document_files(docs, ["backups"], with_scripts)]
    assert rel == ["a.md", "run.sh"]


def test_supported_formats_is_the_union_of_both_tiers():
    assert SUPPORTED_FORMATS == DEFAULT_INDEXED_EXTENSIONS | DEFAULT_REGISTERED_EXTENSIONS
    # Prose stays embedded; code and config are registered (4.4). The code
    # extensions MOVED — they were embedded through 4.3, and the defaults take
    # them out of the embedded list outright rather than relying on precedence.
    assert ".md" in DEFAULT_INDEXED_EXTENSIONS
    assert ".pdf" in DEFAULT_INDEXED_EXTENSIONS
    for ext in (".py", ".sh", ".js", ".css", ".yml", ".yaml", ".toml",
                ".ini", ".sql", ".json5", ".ts", ".cpp"):
        assert ext in DEFAULT_REGISTERED_EXTENSIONS, ext
        assert ext not in DEFAULT_INDEXED_EXTENSIONS, ext
    # .json is the deliberate exception: lorebooks and ST config exports are
    # content, not code, and must stay semantically searchable.
    assert ".json" in DEFAULT_INDEXED_EXTENSIONS
    assert ".json" not in DEFAULT_REGISTERED_EXTENSIONS
    assert not (DEFAULT_INDEXED_EXTENSIONS & DEFAULT_REGISTERED_EXTENSIONS)


def test_policy_tier_routing():
    policy = ExtensionPolicy.build([".md", ".txt"], [".py", ".sh"])
    assert policy.tier_for(".md") == TIER_EMBEDDED
    assert policy.tier_for(".PY") == TIER_REGISTERED  # case-insensitive
    assert policy.tier_for(".zip") is None
    assert policy.all_extensions == {".md", ".txt", ".py", ".sh"}


def test_policy_normalizes_bare_extensions():
    policy = ExtensionPolicy.build(["MD", " .txt "], ["py"])
    assert policy.tier_for(".md") == TIER_EMBEDDED
    assert policy.tier_for(".txt") == TIER_EMBEDDED
    assert policy.tier_for(".py") == TIER_REGISTERED


def test_policy_overlap_resolves_to_registered_and_is_reported():
    """Self-test case 13. Registered MUST win: the extensions this tier exists
    for are already in the embedded allowlist, so "embedded wins" would make
    naming one in registered_extensions a no-op and the feature would ship dead."""
    policy = ExtensionPolicy.build([".md", ".py"], [".py", ".sh"])
    assert policy.tier_for(".py") == TIER_REGISTERED
    assert ".py" not in policy.embedded
    assert policy.conflicts == (".py",)
    assert policy.tier_for(".md") == TIER_EMBEDDED  # non-conflicting unaffected


def test_parse_file_tags_tier_and_skips_chunking_for_registered(docs):
    (docs / "s.py").write_text("def build_worldbook():\n    return 1\n", encoding="utf-8")
    doc = parse_file(docs / "s.py", docs, policy=ExtensionPolicy.build([".md"], [".py"]))
    assert doc.tier == TIER_REGISTERED
    assert doc.is_registered
    assert doc.chunks() == []  # never chunked -> never embedded
    assert "build_worldbook" in doc.content


def test_parse_file_reads_registered_extension_with_no_extractor(docs):
    """.sh/.css/.toml have no dedicated extractor; the registered tier falls
    back to a plain text read rather than refusing the file."""
    (docs / "run.sh").write_text("#!/bin/sh\necho hi\n", encoding="utf-8")
    doc = parse_file(docs / "run.sh", docs, policy=ExtensionPolicy.build([], [".sh"]))
    assert doc.tier == TIER_REGISTERED
    assert "echo hi" in doc.content


def test_parse_file_rejects_extension_outside_both_tiers(docs):
    (docs / "a.md").write_text("x", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported format"):
        parse_file(docs / "a.md", docs, policy=ExtensionPolicy.build([".txt"], [".py"]))


def test_detect_category_longest_match_wins():
    mappings = {"docs": "documentation", "docs/security": "security"}
    assert detect_category("docs/security/notes.md", mappings) == "security"
    assert detect_category("docs/readme.md", mappings) == "documentation"
    assert detect_category("other/x.md", mappings) == "general"
    assert detect_category("anything.md", {}) == "general"


def test_extract_keywords_routes_only():
    routes = {"gpu": ["rocm", "cuda"], "net": ["pcie"]}
    content = "Building ROCm on the new PCIe riser"
    assert extract_keywords(content, routes) == ["pcie", "rocm"]
    assert extract_keywords(content, {}) == []


# ------------------------------------------ 5.1: encodings other than plain UTF-8


@pytest.mark.parametrize("encoding,label", [
    ("utf-16", "utf-16 with BOM (Excel's 'Unicode Text' export)"),
    ("utf-32", "utf-32 with BOM"),
    ("utf-8-sig", "utf-8 with BOM"),
])
def test_read_text_sniffs_the_bom(docs, encoding, label):
    """_read_text assumed UTF-8 with errors="ignore" and no sniff.

    A UTF-16 file decoded that way becomes text with a NUL between every
    character; PostgreSQL text cannot hold U+0000, so replace_document failed
    inside its transaction and the document appeared in summary["errors"] as an
    opaque database error — never indexed, for a reason nobody could read.
    """
    f = docs / "note.md"
    f.write_bytes("# Title\n\nbody text\n".encode(encoding))
    text = _read_text(f)
    assert text == "# Title\n\nbody text\n"
    assert "\x00" not in text
    assert not text.startswith("\ufeff")


def test_read_text_still_normalizes_newlines(docs):
    """Path.read_text did universal-newline translation for free in text mode;
    reading bytes does not. Losing it would put CRLF into the INDEXED copy,
    breaking the frontmatter regex, the section splitter and every anchor built
    from indexed text. The STORED file stays byte-verbatim — that is
    _write_verbatim's job, and a different one."""
    f = docs / "crlf.md"
    f.write_bytes(b"# Title\r\n\r\nbody\r\n")
    assert _read_text(f) == "# Title\n\nbody\n"

    f2 = docs / "cr.md"
    f2.write_bytes(b"old\rmac\r")
    assert _read_text(f2) == "old\nmac\n"
